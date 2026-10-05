import os
import logging
from datetime import datetime
from openpyxl import Workbook, load_workbook

logger = logging.getLogger(__name__)

def log_content_data(content_payload, excel_file="content_log.xlsx"):
    """
    Logs generated script metadata into an Excel file.
    """
    try:
        if os.path.exists(excel_file):
            wb = load_workbook(excel_file)
            ws = wb.active
        else:
            wb = Workbook()
            ws = wb.active
            ws.append(["Timestamp", "Title", "Script", "Tags"])

        timestamp = datetime.now().strftime("%Y-%m-%d %H:%M:%S")
        title = content_payload.get("title", "")
        script = content_payload.get("script", "")
        tags = ", ".join(content_payload.get("tags", []))

        ws.append([timestamp, title, script, tags])
        wb.save(excel_file)
        logger.info(f"Successfully logged entry to {excel_file}")
    except Exception as e:
        logger.error(f"Error writing to Excel sheet: {e}")